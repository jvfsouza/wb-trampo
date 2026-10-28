"""Gera a planilha Excel de acesso rápido a partir do SQLite."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from .banco import Banco

log = logging.getLogger(__name__)

# coluna do banco -> título na planilha (na ordem em que aparecem)
COLUNAS = {
    "titulo": "Título",
    "empresa": "Empresa",
    "site": "Site",
    "local": "Local",
    "modalidade": "Modalidade",
    "nivel": "Nível",
    "tipo_contrato": "Contrato",
    "area": "Área",
    "publicada_em": "Publicada em",
    "candidatos": "Candidatos",
    "url": "Link",
    "busca": "Busca",
    "primeiro_visto_em": "Visto pela 1ª vez",
    "descricao": "Descrição",
}
LIMITE_DESCRICAO = 500
LARGURA_MAX = 60


def _carregar(banco: Banco, filtro: str | None) -> tuple[pd.DataFrame, pd.DataFrame]:
    todas = pd.read_sql(
        f"SELECT {', '.join(COLUNAS)}, execucao FROM vagas "
        "ORDER BY primeiro_visto_em DESC, publicada_em DESC",
        banco.con,
    )
    if filtro:
        texto = todas["titulo"].fillna("") + " " + todas["descricao"].fillna("")
        todas = todas[texto.str.contains(filtro, case=False, regex=False)]

    ultima = banco.ultima_execucao()
    novas = todas[todas["execucao"] == ultima]

    def preparar(df: pd.DataFrame) -> pd.DataFrame:
        df = df.drop(columns="execucao").copy()
        desc = df["descricao"].fillna("")
        cortada = desc.str.len() > LIMITE_DESCRICAO
        df["descricao"] = desc.where(~cortada, desc.str[:LIMITE_DESCRICAO] + "…")
        return df.rename(columns=COLUNAS)

    return preparar(novas), preparar(todas)


def _formatar(ws, df: pd.DataFrame) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        if coluna == "Descrição":
            largura = 80
        else:
            maiores = [len(str(v)) for v in df[coluna].dropna().head(200)]
            largura = min(max([len(coluna), *maiores]) + 2, LARGURA_MAX)
        ws.column_dimensions[letra].width = largura

        if coluna == "Link":
            for celula in ws[letra][1:]:
                if celula.value:
                    celula.hyperlink = celula.value
                    celula.style = "Hyperlink"


def _escrever(caminho: Path, abas: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        for nome, df in abas.items():
            df.to_excel(writer, sheet_name=nome, index=False)
            _formatar(writer.sheets[nome], df)


def exportar(
    banco: Banco,
    saida: str | Path = "data/vagas.xlsx",
    filtro: str | None = None,
) -> Path:
    """Escreve a planilha e retorna o caminho usado.

    Se o arquivo estiver aberto no Excel (o Windows bloqueia a escrita),
    salva com data e hora no nome em vez de falhar.
    """
    saida = Path(saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    log.info("Gerando planilha%s...", f" (filtro: '{filtro}')" if filtro else "")
    novas, todas = _carregar(banco, filtro)
    abas = {"Novas": novas, "Todas": todas}
    try:
        _escrever(saida, abas)
    except PermissionError:
        alternativo = saida.with_name(f"{saida.stem}_{datetime.now():%Y-%m-%d_%H%M}{saida.suffix}")
        log.warning("%s está aberto (provavelmente no Excel); salvando em %s.", saida, alternativo)
        _escrever(alternativo, abas)
        saida = alternativo
    log.info("Planilha salva: %s (%d novas, %d no total)", saida, len(novas), len(todas))
    return saida
