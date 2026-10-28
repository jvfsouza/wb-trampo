"""Fluxo completo sem internet: coleta com cliente falso -> SQLite -> Excel."""

from openpyxl import load_workbook

from wb_trampo.banco import Banco
from wb_trampo.config import Busca, Config, Limites
from wb_trampo.exportar import exportar
from wb_trampo.scraper import executar


class ClienteFalso:
    """Devolve a lista de exemplo na 1ª página, vazio nas seguintes, e o detalhe de exemplo."""

    def __init__(self, html_lista, html_detalhe):
        self.html_lista = html_lista
        self.html_detalhe = html_detalhe
        self.urls = []

    def get(self, url):
        self.urls.append(url)
        if "jobPosting/" in url:
            return self.html_detalhe
        return self.html_lista if "start=0" in url else ""


def _config(**limites):
    return Config(buscas=[Busca("python", "Brasil", modalidade=["remoto"])], limites=Limites(**limites))


def test_coleta_salva_e_nao_duplica(tmp_path, html_lista, html_detalhe):
    banco = Banco(tmp_path / "vagas.db")
    cliente = ClienteFalso(html_lista, html_detalhe)

    r1 = executar(_config(), banco, cliente)
    assert (r1.vistas, r1.novas) == (3, 3)
    linha = banco.con.execute(
        "SELECT url, nivel, modalidade, busca FROM vagas WHERE job_id = '4012345678'"
    ).fetchone()
    assert linha == ("https://www.linkedin.com/jobs/view/4012345678/", "Assistente", "remoto", "python")

    r2 = executar(_config(), banco, cliente)
    assert (r2.vistas, r2.novas) == (3, 0)
    assert banco.contar() == 3

    # 2ª execução não teve vagas novas -> aba "Novas" vazia, "Todas" com 3
    wb = load_workbook(exportar(banco, tmp_path / "vagas.xlsx"))
    assert wb["Novas"].max_row == 1
    assert wb["Todas"].max_row == 4


def test_respeita_limite_e_sem_descricao(tmp_path, html_lista, html_detalhe):
    banco = Banco(tmp_path / "vagas.db")
    cliente = ClienteFalso(html_lista, html_detalhe)
    r = executar(_config(max_vagas_por_busca=2, buscar_descricao=False), banco, cliente)
    assert (r.vistas, r.novas) == (2, 2)
    assert not any("jobPosting/" in u for u in cliente.urls)


def test_excel_tem_abas_e_links(tmp_path, html_lista, html_detalhe):
    banco = Banco(tmp_path / "vagas.db")
    executar(_config(), banco, ClienteFalso(html_lista, html_detalhe))
    wb = load_workbook(exportar(banco, tmp_path / "vagas.xlsx"))

    assert wb.sheetnames == ["Novas", "Todas"]
    ws = wb["Novas"]
    cabecalho = [c.value for c in ws[1]]
    assert ws.max_row == 4
    col_link = cabecalho.index("Link") + 1
    celula = ws.cell(row=2, column=col_link)
    assert celula.value.startswith("https://www.linkedin.com/jobs/view/")
    assert celula.hyperlink.target == celula.value
    assert ws.freeze_panes == "A2"


def test_excel_filtro(tmp_path, html_lista, html_detalhe):
    banco = Banco(tmp_path / "vagas.db")
    executar(_config(), banco, ClienteFalso(html_lista, html_detalhe))
    wb = load_workbook(exportar(banco, tmp_path / "f.xlsx", filtro="dados"))
    assert wb["Todas"].max_row == 2  # só "Engenheira(o) de Dados Pleno"


def test_linkedin_e_gupy_na_mesma_execucao(tmp_path, html_lista, html_detalhe):
    import json
    from pathlib import Path

    resposta_gupy = json.loads((Path(__file__).parent / "fixtures" / "gupy" / "jobs.json").read_text(encoding="utf-8"))

    class ClienteMisto(ClienteFalso):
        def get_json(self, url):
            self.urls.append(url)
            return resposta_gupy if "offset=0" in url else {"data": [], "pagination": {"total": 4}}

    config = Config(
        buscas=[Busca("python", "Brasil"), Busca("backend", site="gupy")],
        limites=Limites(),
    )
    banco = Banco(tmp_path / "vagas.db")
    r = executar(config, banco, ClienteMisto(html_lista, html_detalhe))
    assert (r.vistas, r.novas) == (7, 7)
    por_site = dict(banco.con.execute("SELECT site, COUNT(*) FROM vagas GROUP BY site").fetchall())
    assert por_site == {"linkedin": 3, "gupy": 4}

    wb = load_workbook(exportar(banco, tmp_path / "vagas.xlsx"))
    cabecalho = [c.value for c in wb["Todas"][1]]
    assert cabecalho[:3] == ["Título", "Empresa", "Site"]
    col_link = cabecalho.index("Link") + 1
    links = {wb["Todas"].cell(row=i, column=col_link).value for i in range(2, 9)}
    assert "https://acme.gupy.io/job/eyJqb2JJZCI6MTI3ODYxMjJ9" in links


def test_teto_de_requisicoes_salva_o_que_coletou(tmp_path, html_lista, html_detalhe):
    from wb_trampo.cliente import Cliente

    class Sessao:
        headers = {}

        def get(self, url, timeout):
            class R:
                status_code = 200
                text = html_detalhe if "jobPosting/" in url else html_lista
            return R()

    banco = Banco(tmp_path / "vagas.db")
    config = Config(
        buscas=[Busca("python", "Brasil"), Busca("java", "Brasil")],
        limites=Limites(max_requisicoes=3),
    )
    cliente = Cliente(0, 0, session=Sessao(), dormir=lambda s: None,
                      pausas_por_host={}, max_requisicoes=3)
    r = executar(config, banco, cliente)
    # 1 página + 2 detalhes; o 3º detalhe esbarra no teto e a 2ª busca nem começa
    assert r.interrompida and r.novas == 2
    assert banco.contar() == 2
    assert banco.con.execute("SELECT terminada_em FROM execucoes").fetchone()[0] is not None


def test_intervalo_minimo_entre_execucoes(tmp_path, monkeypatch, caplog):
    import logging

    from wb_trampo import __main__ as cli
    from wb_trampo.scraper import Resumo

    config = tmp_path / "config.toml"
    config.write_text('[[buscas]]\nsite = "gupy"\npalavra_chave = "py"\n', encoding="utf-8")
    chamadas = []
    monkeypatch.setattr(cli, "executar", lambda *a, **k: chamadas.append(1) or Resumo())
    banco = Banco(tmp_path / "vagas.db")
    banco.iniciar_execucao()  # uma execução agora mesmo
    banco.fechar()

    base = ["--banco", str(tmp_path / "vagas.db"), "--logs", str(tmp_path / "logs"),
            "buscar", "--config", str(config), "--sem-excel", "--sem-sheets"]
    with caplog.at_level(logging.INFO):
        assert cli.main(base) == 0
    assert chamadas == [] and "--forcar" in caplog.text

    assert cli.main([*base, "--forcar"]) == 0
    assert chamadas == [1]
